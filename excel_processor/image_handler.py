"""
Image handler module for processing and adding images to Excel sheets.
"""

from typing import List, Tuple, Optional
import os
import re
from openpyxl.drawing.image import Image
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage, ImageOps


class ImageHandler:
    """
    Handles image operations for Excel sheets.
    
    This class provides methods for fixing image orientation, calculating cell dimensions,
    and adding images to Excel sheets.
    """
    
    def __init__(self) -> None:
        """
        Initialize the image handler.
        """
        pass
    
    def fix_image_orientation(self, image_path: str) -> str:
        """
        Fix image orientation based on EXIF data and return path to corrected image.
        Uses PIL's ImageOps.exif_transpose() to automatically handle all EXIF orientations.
        
        Args:
            image_path (str): Path to the original image file
        
        Returns:
            str: Path to the corrected image (may be original if no rotation needed)
        """
        try:
            # Open image with PIL
            pil_img = PILImage.open(image_path)
            
            # Check if image has EXIF orientation data that needs correction
            needs_correction: bool = False
            try:
                exif = pil_img.getexif()
                if exif:
                    orientation = exif.get(274)  # EXIF orientation tag
                    # Orientation values: 1=normal, 2-8 need correction
                    if orientation and orientation != 1:
                        needs_correction = True
            except (AttributeError, KeyError, TypeError):
                # No EXIF data or can't read it
                pass
            
            if needs_correction:
                # Apply EXIF orientation correction automatically
                # This handles all EXIF orientation tags (1-8)
                pil_img_corrected = ImageOps.exif_transpose(pil_img)
                
                # Save to temporary file with same extension as original
                file_ext = os.path.splitext(image_path)[1].lower()
                base_name = os.path.splitext(image_path)[0]
                temp_path = base_name + "_temp_orient" + file_ext
                
                # Convert to RGB if necessary (for JPEG)
                if pil_img_corrected.mode in ("RGBA", "P", "LA"):
                    pil_img_corrected = pil_img_corrected.convert("RGB")
                # Preserve original format if possible
                if file_ext in ['.jpg', '.jpeg']:
                    pil_img_corrected.save(temp_path, quality=95, format='JPEG')
                elif file_ext == '.png':
                    pil_img_corrected.save(temp_path, format='PNG')
                else:
                    pil_img_corrected.save(temp_path, quality=95)
                
                pil_img.close()
                pil_img_corrected.close()
                return temp_path
            
            pil_img.close()
            return image_path
            
        except Exception as e:
            print(f"Warning: Could not fix image orientation for {image_path}: {e}")
            return image_path
    
    def get_cell_dimensions_emu(self, sheet: object, cell_address: str) -> Tuple[int, int]:
        """
        Get cell dimensions in EMU (English Metric Units) for image positioning.
        Handles merged cells by calculating the total merged range dimensions.
        
        Args:
            sheet: The openpyxl worksheet object
            cell_address (str): Cell address (e.g., 'K7', 'AA10')
        
        Returns:
            Tuple[int, int]: (width_emu, height_emu) dimensions in EMU units
        """
        # Parse cell address using regex: column letters followed by row number
        match = re.match(r'^([A-Z]+)(\d+)$', cell_address.upper())
        if not match:
            raise ValueError(f"Invalid cell address: {cell_address}")
        
        col_letter = match.group(1)
        row_num = int(match.group(2))
        
        # Check if the cell is part of a merged range
        merged_range = None
        for merged_cell in sheet.merged_cells.ranges:
            if cell_address in merged_cell:
                merged_range = merged_cell
                break
        
        if merged_range:
            # Calculate dimensions of merged range
            # Get start and end columns and rows
            min_col = merged_range.min_col
            max_col = merged_range.max_col
            min_row = merged_range.min_row
            max_row = merged_range.max_row
            
            # Calculate total width (sum of all columns in merged range)
            total_width = 0.0
            for col_idx in range(min_col, max_col + 1):
                col_letter_merged = get_column_letter(col_idx)
                column_width = sheet.column_dimensions[col_letter_merged].width
                if column_width is None:
                    column_width = 8.43  # Default Excel column width
                total_width += column_width
            
            # Calculate total height (sum of all rows in merged range)
            total_height = 0.0
            for row_idx in range(min_row, max_row + 1):
                row_height = sheet.row_dimensions[row_idx].height
                if row_height is None:
                    row_height = 15  # Default Excel row height
                total_height += row_height
        else:
            # Single cell - use original logic
            column_width = sheet.column_dimensions[col_letter].width
            if column_width is None:
                column_width = 8.43  # Default Excel column width
            total_width = column_width
            
            row_height = sheet.row_dimensions[row_num].height
            if row_height is None:
                row_height = 15  # Default Excel row height
            total_height = row_height
        
        # Convert to pixels and then to EMU
        # Column width: 1 character ≈ 7 pixels at 96 DPI
        # Row height: 1 point = 4/3 pixels at 96 DPI
        # 1 pixel = 9525 EMU
        width_pixels = total_width * 7
        width_emu = int(width_pixels * 9525)
        
        height_pixels = total_height * (4 / 3)
        height_emu = int(height_pixels * 9525)
        
        return (width_emu, height_emu)
    
    def add_image_to_cell(self, sheet: object, img: Image, cell_address: str) -> None:
        """
        Add an image to a cell, resizing it to fit inside the cell.
        Handles merged cells and allows scaling up to fill the cell better.
        
        Args:
            sheet: The openpyxl worksheet object
            img: The Image object to add
            cell_address (str): Cell address where to place the image (e.g., 'K7')
        """
        # Get cell dimensions in EMU (handles merged cells)
        cell_width_emu, cell_height_emu = self.get_cell_dimensions_emu(sheet, cell_address)
        
        # Use 95% of cell dimensions to leave a small margin
        usable_width_emu = int(cell_width_emu * 0.95)
        usable_height_emu = int(cell_height_emu * 0.95)
        
        # Get current image dimensions in EMU (openpyxl uses EMU internally)
        # Image dimensions are already in pixels, convert to EMU
        # 1 pixel = 9525 EMU
        img_width_emu = int(img.width * 9525)
        img_height_emu = int(img.height * 9525)
        
        # Calculate scaling factor to fit image inside cell
        # Allow scaling up if image is smaller than cell
        width_ratio = usable_width_emu / img_width_emu
        height_ratio = usable_height_emu / img_height_emu
        scale_factor = min(width_ratio, height_ratio)  # Use smaller ratio to maintain aspect ratio
        
        # Resize image to fit within cell (can scale up or down)
        img.width = int(img.width * scale_factor)
        img.height = int(img.height * scale_factor)
        print(f"Resized image for cell {cell_address}: {img.width}x{img.height} pixels (scale factor: {scale_factor:.2f})")
        
        # Set anchor to cell (this positions the image at the top-left of the cell)
        img.anchor = cell_address
        
        # Add image to sheet
        sheet.add_image(img)
    
    def add_image_to_sheet(self, sheet: object, sheet_id: str, img_dir: str) -> Tuple[bool, List[str]]:
        """
        Add images to sheet based on sheet ID.
        
        Args:
            sheet: The openpyxl worksheet object
            sheet_id (str): The sheet ID to match with image folder
            img_dir (str): Path to the image directory
        
        Returns:
            Tuple[bool, List[str]]: (success: bool, temp_files: List[str]) - Success status and list of temporary files to clean up later
        """
        temp_files: List[str] = []  # Track temporary files for cleanup
        try:
            images_path = os.path.join(img_dir, sheet_id)
            if not os.path.exists(images_path):
                print(f"Warning: Images not found for {sheet_id}: {images_path}")
                return (False, temp_files)
            images = os.listdir(images_path)
            site_images: List[str] = []
            plan_images: List[str] = []
            for image in images:
                if "site" in image.lower():
                    site_images.append(image)
                elif "plan" in image.lower():
                    plan_images.append(image)
            
            # Add site images
            for i, site_image in enumerate(site_images):
                image_path = os.path.join(images_path, site_image)
                try:
                    # Fix orientation before loading
                    corrected_path = self.fix_image_orientation(image_path)
                    if corrected_path != image_path:
                        temp_files.append(corrected_path)
                    
                    if not os.path.exists(corrected_path):
                        print(f"Error: Image file not found: {corrected_path}")
                        continue
                    
                    img = Image(corrected_path)
                    
                    # Determine cell address based on image index
                    if i == 0:
                        cell_address = 'K7'
                    else:
                        cell_address = 'O7'
                    
                    # Add image to cell (will be resized to fit)
                    self.add_image_to_cell(sheet, img, cell_address)
                    print(f"Added site image {site_image} to cell {cell_address}")
                except Exception as e:
                    print(f"Error adding site image {site_image}: {e}")
            
            # Add plan images
            for i, plan_image in enumerate(plan_images):
                image_path = os.path.join(images_path, plan_image)
                try:
                    # Fix orientation before loading
                    corrected_path = self.fix_image_orientation(image_path)
                    if corrected_path != image_path:
                        temp_files.append(corrected_path)
                    
                    if not os.path.exists(corrected_path):
                        print(f"Error: Image file not found: {corrected_path}")
                        continue
                    
                    img = Image(corrected_path)
                    
                    # Determine cell address based on image index
                    if i == 0:
                        cell_address = 'O16'
                    else:
                        # If there's a second plan image, place it in a different cell
                        # Adjust this cell address as needed
                        cell_address = 'K16'
                    
                    # Add image to cell (will be resized to fit)
                    self.add_image_to_cell(sheet, img, cell_address)
                    print(f"Added plan image {plan_image} to cell {cell_address}")
                except Exception as e:
                    print(f"Error adding plan image {plan_image}: {e}")
            
            print(f"Added images for {sheet_id}")
            return (True, temp_files)
            
        except Exception as e:
            print(f"Error adding image for {sheet_id}: {e}")
            import traceback
            traceback.print_exc()
            return (False, temp_files)


