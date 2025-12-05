"""
Sheet finder module for locating and validating Excel sheets.
"""

from typing import List, Optional, Tuple
from openpyxl import Workbook


class SheetFinder:
    """
    Handles finding and validating Excel sheets.
    
    This class provides methods for finding specific sheets in a workbook
    and validating that required sheets exist.
    """
    
    def __init__(self) -> None:
        """
        Initialize the sheet finder.
        """
        pass
    
    def find_sheets_containing(self, workbook: Workbook, search_term: str) -> List[str]:
        """
        Find all sheet names containing the search term (case-insensitive).
        
        Args:
            workbook: The openpyxl workbook object
            search_term (str): The term to search for in sheet names
        
        Returns:
            List[str]: List of sheet names containing the search term
        """
        matching_sheets: List[str] = []
        for sheet_name in workbook.sheetnames:
            if search_term.lower() in sheet_name.lower():
                matching_sheets.append(sheet_name)
        return matching_sheets
    
    def find_decision_matrix_sheet(self, workbook: Workbook) -> Optional[str]:
        """
        Find the sheet that contains a cell with the exact content "Decision Matrix" (case-insensitive).
        
        Args:
            workbook: The openpyxl workbook object
        
        Returns:
            Optional[str]: The name of the sheet containing "Decision Matrix" in a cell, or None if not found
        """
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            # Search through all cells in the sheet
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cell_value = str(cell.value).strip()
                        if cell_value.lower() == "decision matrix":
                            return sheet_name
        return None
    
    def find_template_sheets(self, workbook: Workbook) -> List[str]:
        """
        Find all sheets containing "Template" in their name.
        
        Args:
            workbook: The openpyxl workbook object
        
        Returns:
            List[str]: List of sheet names containing "Template"
        """
        return self.find_sheets_containing(workbook, "Template")
    
    def validate_required_sheets(self, workbook: Workbook, template_sheet_name: Optional[str] = None) -> Tuple[bool, str]:
        """
        Validate that all required sheets exist in the workbook.
        
        Args:
            workbook: The openpyxl workbook object
            template_sheet_name (Optional[str]): Name of the template sheet to check. If None, will search for it.
        
        Returns:
            Tuple[bool, str]: (is_valid, error_message) - True if all sheets exist, False with error message otherwise
        """
        missing_sheets: List[str] = []
        available_sheets = workbook.sheetnames
        
        # Check for Cover sheet
        cover_found: bool = False
        for sheet_name in available_sheets:
            if sheet_name.lower() == "cover":
                cover_found = True
                break
        if not cover_found:
            missing_sheets.append("Cover")
        
        # Check for GenInfo+Contacts sheet
        geninfo_found: bool = False
        for sheet_name in available_sheets:
            if sheet_name.lower() == "geninfo+contacts":
                geninfo_found = True
                break
        if not geninfo_found:
            missing_sheets.append("GenInfo+Contacts")
        
        # Check for Decision Matrix sheet
        decision_matrix_sheet = self.find_decision_matrix_sheet(workbook)
        if not decision_matrix_sheet:
            missing_sheets.append("Decision Matrix (sheet containing a cell with 'Decision Matrix')")
        
        # Check for Template sheet
        if template_sheet_name:
            if template_sheet_name not in available_sheets:
                missing_sheets.append(f"Template sheet '{template_sheet_name}'")
        else:
            template_sheets = self.find_template_sheets(workbook)
            if not template_sheets:
                missing_sheets.append("Template (sheet containing 'Template' in name)")
        
        if missing_sheets:
            error_msg = "The following required sheets are missing:\n"
            error_msg += "\n".join(f"  - {sheet}" for sheet in missing_sheets)
            return (False, error_msg)
        
        return (True, "")
    
    def find_image_cell_in_cover_sheet(self, workbook: Workbook) -> Optional[str]:
        """
        Find the cell address containing the term "image" (case-insensitive) in the Cover sheet.
        
        Args:
            workbook: The openpyxl workbook object
        
        Returns:
            Optional[str]: Cell address (e.g., 'A1') containing "image", or None if not found
        """
        cover_sheet_name: Optional[str] = None
        for sheet_name in workbook.sheetnames:
            if sheet_name.lower() == "cover":
                cover_sheet_name = sheet_name
                break
        
        if not cover_sheet_name:
            return None
        
        sheet = workbook[cover_sheet_name]
        # Search through all cells in the sheet
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cell_value = str(cell.value).strip()
                    if "image" in cell_value.lower():
                        # Return cell address in Excel format (e.g., 'A1')
                        from openpyxl.utils import get_column_letter
                        cell_address = f"{get_column_letter(cell.column)}{cell.row}"
                        return cell_address
        return None


