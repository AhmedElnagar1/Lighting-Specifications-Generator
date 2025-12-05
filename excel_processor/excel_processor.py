"""
Main Excel processor module that orchestrates the processing workflow.
"""

from typing import Optional, Union
import os
import shutil
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from excel_processor.image_handler import ImageHandler
from excel_processor.sheet_finder import SheetFinder
from excel_processor.sheet_creator import SheetCreator
from excel_processor.pdf_exporter import PDFExporter
from excel_processor.debug_logger import DebugLogger
from excel_processor.data_mapper import DataMapper


class ExcelProcessor:
    """
    Main orchestrator class for processing Excel files.
    
    This class coordinates all the components needed to process Excel files,
    create sheets, map data, and generate PDFs.
    """
    
    def __init__(self) -> None:
        """
        Initialize the Excel processor with all required components.
        """
        self.debug_logger = DebugLogger()
        self.image_handler = ImageHandler()
        self.sheet_finder = SheetFinder()
        self.data_mapper = DataMapper(self.debug_logger)
        self.sheet_creator = SheetCreator(
            image_handler=self.image_handler,
            data_mapper=self.data_mapper,
            debug_logger=self.debug_logger
        )
        self.pdf_exporter = PDFExporter()
    
    def create_backup(self, excel_file_path: str) -> Optional[str]:
        """
        Create a backup copy of the original file in a Backup folder.
        
        Args:
            excel_file_path (str): Path to the Excel file to backup
        
        Returns:
            Optional[str]: Path to the backup file, or None if backup failed
        """
        try:
            # Get the directory of the Excel file
            file_dir = os.path.dirname(os.path.abspath(excel_file_path))
            
            # Create Backup folder path
            backup_dir = os.path.join(file_dir, "Backup")
            
            # Create Backup folder if it doesn't exist
            os.makedirs(backup_dir, exist_ok=True)
            
            # Get the base filename without extension
            base_name = os.path.splitext(os.path.basename(excel_file_path))[0]
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            
            # Create backup filename
            backup_filename = f"{base_name}_backup_{timestamp}.xlsx"
            backup_path = os.path.join(backup_dir, backup_filename)
            
            # Copy the file to the backup location
            shutil.copy2(excel_file_path, backup_path)
            print(f"Created backup: {backup_path}")
            return backup_path
        except Exception as e:
            print(f"Warning: Could not create backup: {e}")
            return None
    
    def add_cover_image(self, excel_file_path: str, image_path: str) -> bool:
        """
        Add an image to the Cover sheet at the cell containing the term "image".
        
        Args:
            excel_file_path (str): Path to the Excel file
            image_path (str): Path to the image file to add
        
        Returns:
            bool: True if successful, False otherwise
        """
        try:
            # Load workbook
            workbook = load_workbook(excel_file_path)
            
            # Find the cell containing "image" in Cover sheet
            cell_address = self.sheet_finder.find_image_cell_in_cover_sheet(workbook)
            if not cell_address:
                print("Error: Could not find cell containing 'image' in Cover sheet")
                workbook.close()
                return False
            
            # Get the Cover sheet
            cover_sheet_name: Optional[str] = None
            for sheet_name in workbook.sheetnames:
                if sheet_name.lower() == "cover":
                    cover_sheet_name = sheet_name
                    break
            
            if not cover_sheet_name:
                print("Error: Cover sheet not found")
                workbook.close()
                return False
            
            sheet = workbook[cover_sheet_name]
            
            # Remove existing images in the target cell area (if any)
            # Check all images and remove those anchored to the target cell
            images_to_remove = []
            for img in list(sheet._images):  # Create a copy of the list to iterate safely
                if hasattr(img, 'anchor') and img.anchor:
                    # Check if image anchor matches our target cell
                    anchor_str = str(img.anchor)
                    # The anchor might be a cell reference like 'A1' or a range like 'A1:A1'
                    # Check if the cell address is part of the anchor
                    if cell_address in anchor_str:
                        images_to_remove.append(img)
            
            # Remove old images
            for img in images_to_remove:
                try:
                    sheet._images.remove(img)
                except ValueError:
                    # Image might have already been removed
                    pass
            
            # Fix image orientation if needed
            corrected_path = self.image_handler.fix_image_orientation(image_path)
            temp_files = []
            if corrected_path != image_path:
                temp_files.append(corrected_path)
            
            # Load and add the image
            img = Image(corrected_path)
            self.image_handler.add_image_to_cell(sheet, img, cell_address)
            
            # Save the workbook
            workbook.save(excel_file_path)
            workbook.close()
            
            # Clean up temporary files
            for temp_file in temp_files:
                try:
                    if os.path.exists(temp_file):
                        os.remove(temp_file)
                except Exception as e:
                    print(f"Warning: Could not delete temporary file {temp_file}: {e}")
            
            print(f"Successfully added cover image to cell {cell_address}")
            return True
            
        except Exception as e:
            print(f"Error adding cover image: {e}")
            import traceback
            traceback.print_exc()
            return False
    
    def process_excel_file(
        self,
        input_file: str,
        img_dir: str,
        template_sheet_name: Optional[str] = None
    ) -> Union[str, str]:
        """
        Main processing function.
        
        Args:
            input_file (str): Path to the input Excel file
            img_dir (str): Path to the image directory
            template_sheet_name (Optional[str]): Name of the template sheet to use. If None, will be auto-detected.
        
        Returns:
            Union[str, str]: Path to the generated PDF file on success, or error message string on error
        """
        print("Starting Excel processing and PDF creation...")
        print("=" * 50)
        
        # Create backup
        backup_path = self.create_backup(input_file)
        
        # Load workbook
        try:
            input_wb = load_workbook(input_file, data_only=True)  # Read only values, not formulas
            print(f"Loaded workbook: {input_file}")
        except Exception as e:
            error_msg = f"Error loading workbook: {str(e)}"
            print(error_msg)
            return error_msg
        
        if input_wb is None:
            return "Error: Failed to load workbook"
        
        # Validate all required sheets exist
        is_valid, validation_error = self.sheet_finder.validate_required_sheets(input_wb, template_sheet_name)
        if not is_valid:
            print(validation_error)
            input_wb.close()
            return validation_error
        
        # Find Decision Matrix sheet (sheet containing a cell with "Decision Matrix")
        decision_matrix_sheet = self.sheet_finder.find_decision_matrix_sheet(input_wb)
        if not decision_matrix_sheet:
            error_msg = "No sheet containing a cell with 'Decision Matrix' found"
            print(f"Error: {error_msg}")
            input_wb.close()
            return error_msg
        print(f"Found Decision Matrix sheet: {decision_matrix_sheet}")
        
        # Find or use provided template sheet
        if template_sheet_name is None:
            template_sheets = self.sheet_finder.find_template_sheets(input_wb)
            if not template_sheets:
                error_msg = "No sheet containing 'Template' found"
                print(f"Error: {error_msg}")
                input_wb.close()
                return error_msg
            elif len(template_sheets) == 1:
                template_sheet_name = template_sheets[0]
                print(f"Found Template sheet: {template_sheet_name}")
            else:
                # Multiple template sheets found - this should be handled by the GUI
                error_msg = f"Multiple template sheets found: {', '.join(template_sheets)}. Please select a template sheet in the GUI."
                print(f"Error: {error_msg}")
                input_wb.close()
                return error_msg
        else:
            if template_sheet_name not in input_wb.sheetnames:
                error_msg = f"Specified template sheet '{template_sheet_name}' not found"
                print(f"Error: {error_msg}")
                input_wb.close()
                return error_msg
            print(f"Using Template sheet: {template_sheet_name}")
        
        # Determine sheets to keep
        sheets_to_keep = ["Cover", "GenInfo+Contacts", template_sheet_name, decision_matrix_sheet]
        for sheet in input_wb.sheetnames:
            if sheet not in sheets_to_keep:
                input_wb.remove(input_wb[sheet])
        
        # Create sheets
        sheet_ids = self.sheet_creator.create_sheets(
            input_wb, img_dir, input_file, template_sheet_name, decision_matrix_sheet
        )
        if isinstance(sheet_ids, str):  # Error message returned
            input_wb.close()
            return sheet_ids
        
        # Create PDF
        pdf_path = self.pdf_exporter.create_pdf(input_file, sheet_ids)
        if isinstance(pdf_path, str) and not pdf_path.endswith('.pdf'):  # Error message returned
            input_wb.close()
            return pdf_path
        
        input_wb.close()
        
        print("=" * 50)
        print("Processing completed successfully!")
        print(f"Modified Excel file: {input_file}")
        print(f"PDF output: {os.path.splitext(input_file)[0]}_output.pdf")
        if backup_path:
            print(f"Backup created: {backup_path}")
        return pdf_path


# Module-level instance for backward compatibility
_processor = ExcelProcessor()


# Backward compatibility functions
def process_excel_file(
    input_file: str,
    img_dir: str,
    template_sheet_name: Optional[str] = None
) -> Union[str, str]:
    """
    Main processing function (backward compatibility wrapper).
    
    Args:
        input_file (str): Path to the input Excel file
        img_dir (str): Path to the image directory
        template_sheet_name (Optional[str]): Name of the template sheet to use. If None, will be auto-detected.
    
    Returns:
        Union[str, str]: Path to the generated PDF file on success, or error message string on error
    """
    return _processor.process_excel_file(input_file, img_dir, template_sheet_name)


def find_template_sheets(workbook: object) -> list:
    """
    Find all sheets containing "Template" in their name (backward compatibility wrapper).
    
    Args:
        workbook: The openpyxl workbook object
    
    Returns:
        list: List of sheet names containing "Template"
    """
    finder = SheetFinder()
    return finder.find_template_sheets(workbook)


def find_decision_matrix_sheet(workbook: object) -> Optional[str]:
    """
    Find the sheet that contains a cell with the exact content "Decision Matrix" (backward compatibility wrapper).
    
    Args:
        workbook: The openpyxl workbook object
    
    Returns:
        Optional[str]: The name of the sheet containing "Decision Matrix" in a cell, or None if not found
    """
    finder = SheetFinder()
    return finder.find_decision_matrix_sheet(workbook)


def add_cover_image(excel_file_path: str, image_path: str) -> bool:
    """
    Add an image to the Cover sheet at the cell containing the term "image" (backward compatibility wrapper).
    
    Args:
        excel_file_path (str): Path to the Excel file
        image_path (str): Path to the image file to add
    
    Returns:
        bool: True if successful, False otherwise
    """
    return _processor.add_cover_image(excel_file_path, image_path)


def create_backup(excel_file_path: str) -> Optional[str]:
    """
    Create a backup copy of the original file in a Backup folder (backward compatibility wrapper).
    
    Args:
        excel_file_path (str): Path to the Excel file to backup
    
    Returns:
        Optional[str]: Path to the backup file, or None if backup failed
    """
    return _processor.create_backup(excel_file_path)


def add_page_footer(sheet: object, page_number: int, total_pages: int) -> None:
    """
    Add a footer with creation date on the left and page numbering on the right (backward compatibility wrapper).
    
    Args:
        sheet: The openpyxl worksheet object
        page_number (int): Current page number
        total_pages (int): Total number of pages
    """
    creator = SheetCreator()
    creator.add_page_footer(sheet, page_number, total_pages)


