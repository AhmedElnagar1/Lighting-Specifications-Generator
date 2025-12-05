"""
Unit tests for the SheetCreator class.
"""

import pytest
from unittest.mock import MagicMock, PropertyMock
from openpyxl import Workbook
from openpyxl.styles import Font
from excel_processor.sheet_creator import SheetCreator


class TestExtractBoldWordsFromTemplate:
    """Tests for the extract_bold_words_from_template method."""
    
    def test_extracts_bold_words_correctly(self) -> None:
        """Test that bold words are correctly extracted from template sheet."""
        # Create a workbook with some bold and non-bold cells
        wb = Workbook()
        sheet = wb.active
        
        # Add cells with bold formatting
        sheet["A1"] = "Report Code"
        sheet["A1"].font = Font(bold=True)
        
        sheet["B1"] = "Description"
        sheet["B1"].font = Font(bold=True)
        
        # Add non-bold cell
        sheet["C1"] = "Regular Text"
        sheet["C1"].font = Font(bold=False)
        
        # Add another bold cell
        sheet["A2"] = "Manufacturer"
        sheet["A2"].font = Font(bold=True)
        
        creator = SheetCreator()
        bold_words = creator.extract_bold_words_from_template(sheet)
        
        assert "Report Code" in bold_words
        assert "Description" in bold_words
        assert "Manufacturer" in bold_words
        assert "Regular Text" not in bold_words
        assert len(bold_words) == 3
        
        wb.close()
    
    def test_returns_empty_list_for_no_bold_words(self) -> None:
        """Test that empty list is returned when no bold words exist."""
        wb = Workbook()
        sheet = wb.active
        
        # Add only non-bold cells
        sheet["A1"] = "Text 1"
        sheet["A1"].font = Font(bold=False)
        
        sheet["B1"] = "Text 2"
        sheet["B1"].font = Font(bold=False)
        
        creator = SheetCreator()
        bold_words = creator.extract_bold_words_from_template(sheet)
        
        assert bold_words == []
        
        wb.close()


class TestFindHeaderRowByBoldWords:
    """Tests for the find_header_row_by_bold_words method."""
    
    def test_finds_header_row_with_matching_words(self) -> None:
        """Test that header row is found when enough bold words match."""
        wb = Workbook()
        sheet = wb.active
        
        # Row 1 - not enough matches
        sheet["A1"] = "Some text"
        sheet["B1"] = "Other text"
        
        # Row 5 - header row with 7+ matches
        bold_words = ["Report Code", "Description", "Manufacturer", "Power", "Voltage", "Color", "Size", "Price"]
        for col_idx, word in enumerate(bold_words, start=1):
            sheet.cell(row=5, column=col_idx, value=word)
        
        # Row 6 - data row
        sheet["A6"] = "LC-01"
        
        creator = SheetCreator()
        header_row = creator.find_header_row_by_bold_words(sheet, bold_words, min_matches=7)
        
        assert header_row == 5
        
        wb.close()
    
    def test_returns_none_when_no_matching_row(self) -> None:
        """Test that None is returned when no row has enough matches."""
        wb = Workbook()
        sheet = wb.active
        
        # Add some rows with few matches
        sheet["A1"] = "Report Code"
        sheet["B1"] = "Description"
        sheet["C1"] = "Random"
        
        bold_words = ["Report Code", "Description", "Manufacturer", "Power", "Voltage", "Color", "Size", "Price"]
        
        creator = SheetCreator()
        header_row = creator.find_header_row_by_bold_words(sheet, bold_words, min_matches=7)
        
        assert header_row is None
        
        wb.close()

