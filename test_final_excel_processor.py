"""Unit tests for final_excel_processor module."""
from openpyxl import Workbook
from excel_processor.debug_logger import DebugLogger
from excel_processor.sheet_finder import SheetFinder
from excel_processor.data_mapper import DataMapper


def test_debug_logger() -> None:
    """
    Test the DebugLogger class basic functionality.
    """
    # Test 1: Write and close debug logger
    logger = DebugLogger()
    logger.write("Test message 1")
    logger.write("Test message 2")
    logger.close()
    
    # Verify logger can be closed and reopened
    logger2 = DebugLogger()
    logger2.write("Test message 3")
    logger2.close()
    
    assert True, "DebugLogger should work without errors"


def test_sheet_finder() -> None:
    """
    Test the SheetFinder class with a sample workbook.
    """
    # Test 1: Find sheets containing a term
    wb = Workbook()
    wb.create_sheet("Template Sheet 1")
    wb.create_sheet("Template Sheet 2")
    wb.create_sheet("Cover")
    wb.create_sheet("GenInfo+Contacts")
    
    finder = SheetFinder()
    template_sheets = finder.find_template_sheets(wb)
    
    assert len(template_sheets) == 2, "Should find 2 template sheets"
    assert "Template Sheet 1" in template_sheets, "Should find Template Sheet 1"
    assert "Template Sheet 2" in template_sheets, "Should find Template Sheet 2"
    
    # Test 2: Validate required sheets
    is_valid, error_msg = finder.validate_required_sheets(wb)
    assert not is_valid, "Should be invalid (missing Decision Matrix)"
    assert "Decision Matrix" in error_msg, "Error message should mention Decision Matrix"
    
    wb.close()


def test_data_mapper_clean_keyword() -> None:
    """
    Test the DataMapper _clean_keyword method.
    """
    mapper = DataMapper()
    
    # Test 1: Clean keyword with special characters
    keyword1 = "Test*Keyword\nWith\nNewlines"
    cleaned1 = mapper._clean_keyword(keyword1)
    assert "*" not in cleaned1, "Should remove asterisks"
    assert "\n" not in cleaned1, "Should remove newlines"
    
    # Test 2: Clean keyword with extra whitespace
    keyword2 = "  Test   Keyword  "
    cleaned2 = mapper._clean_keyword(keyword2)
    assert cleaned2 == "Test Keyword", "Should normalize whitespace"
    
    # Test 3: Exact match
    match1 = mapper._find_exact_match("Test Keyword", "Test Keyword")
    assert match1 is True, "Should find exact match"
    
    match2 = mapper._find_exact_match("Test Keyword", "test keyword")
    assert match2 is True, "Should find case-insensitive match"
    
    match3 = mapper._find_exact_match("Test Keyword", "Different Text")
    assert match3 is False, "Should not match different text"


if __name__ == "__main__":
    test_debug_logger()
    print("DebugLogger tests passed!")
    
    test_sheet_finder()
    print("SheetFinder tests passed!")
    
    test_data_mapper_clean_keyword()
    print("DataMapper tests passed!")
    
    print("All tests passed!")
